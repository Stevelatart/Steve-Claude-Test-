#!/usr/bin/env python3
"""Render the aggregated DSP metrics (data/dsp-metrics.json) into a formatted
Excel workbook that mirrors the dashboard: an overview, per-dimension sheets,
and the full audience list with Excel AutoFilter (the explorer, in a workbook).

Usage:
    python build_excel.py data/dsp-metrics.json Amazon_DSP_Report_June_2026.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DISPLAY = {"Furnitre Distributors": "Furniture Distributors"}
CLS = {"BEHAVIOR": "In-market / Lifestyle", "USER_AGENT": "Device / OS",
       "CUSTOM_ASIN": "Custom ASIN", "INVENTORY": "Inventory", "OTHER": "Other"}

INK = "16202E"; AMBER = "E8890C"; PAPER = "F1F3F6"; GREEN = "2F855A"
INT_FMT = "#,##0"; PCT_FMT = "0.00%"

thin = Side(style="thin", color="D9DEE5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=INK)
BAND_FILL = PatternFill("solid", fgColor="F7F8FA")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=INK, size=16)
SUB_FONT = Font(name="Calibri", color="7A8797", size=10)
LABEL_FONT = Font(name="Calibri", bold=True, color=INK, size=11)


def name_of(c):
    return DISPLAY.get(c, c)


def style_header(ws, row, ncols):
    for j in range(1, ncols + 1):
        cell = ws.cell(row=row, column=j)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="left" if j == 1 else "right",
                                   vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 26


def write_table(ws, start_row, headers, rows, fmts, band=True):
    """Generic table writer; fmts is a list of number formats per column ('' = text)."""
    ncols = len(headers)
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, ncols)
    r = start_row
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        for j, (val, fmt) in enumerate(zip(row, fmts), 1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if j == 1 else "right", vertical="center")
            if fmt:
                cell.number_format = fmt
            if band and i % 2 == 1:
                cell.fill = BAND_FILL
    return r  # last data row


def autosize(ws, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def build(data, out_path):
    wb = Workbook()
    clients = data["clients"]
    order = data["order"]

    # ---------- Overview ----------
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Amazon DSP — Performance Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Brandography  ·  Reporting period {data['period']}  ·  Generated {data['generated']}"
    ws["A2"].font = SUB_FONT

    hdr = ["Client", "Impressions", "Clicks", "CTR", "Conversions", "Segments", "Sites & apps"]
    rows, fmts = [], ["", INT_FMT, INT_FMT, PCT_FMT, INT_FMT, INT_FMT, INT_FMT]
    for c in order:
        k = clients[c]["kpi"]
        rows.append([name_of(c), k["impr"], k["clicks"], k["ctr"] / 100,
                     k["conv"] if k["conv"] else None, k["segments"], k["sites"]])
    # totals row (CTR = weighted)
    tot_impr = sum(clients[c]["kpi"]["impr"] for c in order)
    tot_clk = sum(clients[c]["kpi"]["clicks"] for c in order)
    rows.append(["All clients", tot_impr, tot_clk, (tot_clk / tot_impr) if tot_impr else 0,
                 sum(clients[c]["kpi"]["conv"] for c in order), None, None])
    last = write_table(ws, 4, hdr, rows, fmts, band=True)
    # bold the totals row
    for j in range(1, len(hdr) + 1):
        ws.cell(row=last, column=j).font = Font(bold=True, color=INK)
        ws.cell(row=last, column=j).fill = PatternFill("solid", fgColor="FCE9CF")
    autosize(ws, [26, 14, 11, 9, 13, 11, 12])
    ws.freeze_panes = "A5"

    note = ("Methodology — Delivery figures (impressions, clicks, CTR, conversions), campaigns, supply "
            "sources and sites come from the DSP inventory report: each impression is counted once against "
            "the site it served on, so these are true de-duplicated totals. The 'Audiences' sheet comes from "
            "the audience-segment report and ranks relative audience performance only — one impression can "
            "match many segments at once, so segment impressions overlap and are NOT summed into delivery "
            "totals. Conversions are off-Amazon (pixel) actions. CTR = clicks ÷ impressions.")
    nr = last + 2
    ws.cell(row=nr, column=1, value=note)
    ws.cell(row=nr, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=nr, column=1).font = SUB_FONT
    ws.merge_cells(start_row=nr, start_column=1, end_row=nr + 5, end_column=7)

    # ---------- Campaigns ----------
    ws = wb.create_sheet("Campaigns")
    ws.sheet_view.showGridLines = False
    rows = []
    for c in order:
        for cm in clients[c]["campaigns"]:
            if cm["impr"] <= 0:
                continue
            rows.append([name_of(c), cm["name"], cm["impr"], cm["clicks"],
                         cm["ctr"] / 100, cm["conv"] if cm["conv"] else None])
    last = write_table(ws, 1, ["Client", "Campaign", "Impressions", "Clicks", "CTR", "Conversions"],
                       rows, ["", "", INT_FMT, INT_FMT, PCT_FMT, INT_FMT])
    autosize(ws, [22, 44, 14, 11, 9, 13])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{last}"

    # ---------- Supply sources ----------
    ws = wb.create_sheet("Supply Sources")
    ws.sheet_view.showGridLines = False
    rows = []
    for c in order:
        for s in clients[c]["supply"]:
            rows.append([name_of(c), s["name"], s["impr"], s["clicks"],
                         s["ctr"] / 100, s["share"] / 100])
    last = write_table(ws, 1, ["Client", "Supply source", "Impressions", "Clicks", "CTR", "Share of impr."],
                       rows, ["", "", INT_FMT, INT_FMT, PCT_FMT, PCT_FMT])
    autosize(ws, [22, 34, 14, 11, 9, 14])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{last}"

    # ---------- Top sites ----------
    ws = wb.create_sheet("Top Sites")
    ws.sheet_view.showGridLines = False
    rows = []
    for c in order:
        for s in clients[c]["sites"]:
            rows.append([name_of(c), s["name"], s["impr"], s["clicks"], s["ctr"] / 100])
    last = write_table(ws, 1, ["Client", "Site or app", "Impressions", "Clicks", "CTR"],
                       rows, ["", "", INT_FMT, INT_FMT, PCT_FMT])
    autosize(ws, [22, 52, 14, 11, 9])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{last}"

    # ---------- Audiences (full, filterable) ----------
    ws = wb.create_sheet("Audiences")
    ws.sheet_view.showGridLines = False
    rows = []
    for c in order:
        for s in clients[c]["segments_all"]:
            rows.append([name_of(c), s["name"], CLS.get(s["cls"], s["cls"]),
                         s["impr"], s["ctr"] / 100, s["conv"] if s["conv"] else None])
    last = write_table(ws, 1, ["Client", "Audience segment", "Type", "Impressions", "CTR", "Conversions"],
                       rows, ["", "", "", INT_FMT, PCT_FMT, INT_FMT], band=False)
    autosize(ws, [22, 52, 20, 14, 9, 13])
    ws.freeze_panes = "C2"       # keep Client + Segment visible while scrolling
    ws.auto_filter.ref = f"A1:F{last}"

    wb.save(out_path)
    return {"clients": len(order), "audience_rows": len(rows)}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: python build_excel.py data/dsp-metrics.json OUTPUT.xlsx")
    info = build(json.load(open(sys.argv[1])), sys.argv[2])
    print(f"wrote {sys.argv[2]}: {info['clients']} clients, {info['audience_rows']} audience rows")
